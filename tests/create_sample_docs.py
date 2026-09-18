"""
Generates realistic sample documents in PDF, DOCX, MD, and TXT formats for Phase 0 testing.
Matches the questions and assertions in tests/test_questions.json.
"""
import os
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
import docx

SAMPLE_DIR = os.path.join(os.path.dirname(__file__), "sample_docs")
os.makedirs(SAMPLE_DIR, exist_ok=True)

def generate_pdf():
    pdf_path = os.path.join(SAMPLE_DIR, "acme_hr_policy.pdf")
    c = canvas.Canvas(pdf_path, pagesize=letter)
    
    # Page 1
    c.setFont("Helvetica-Bold", 16)
    c.drawString(50, 750, "Acme Corporation — Employee Handbook & HR Policy")
    c.setFont("Helvetica-Bold", 12)
    c.drawString(50, 720, "1. Annual Paid Time Off (PTO)")
    c.setFont("Helvetica", 10)
    c.drawString(50, 700, "All full-time employees accrue 20 business days of paid time off (PTO) annually.")
    c.drawString(50, 685, "PTO begins accruing on the first calendar day of employment.")
    c.drawString(50, 670, "Employees may carry over up to 5 unused PTO days into the following calendar year.")
    c.drawString(50, 655, "Any additional accrued days beyond 5 will expire on December 31st without compensation.")
    
    c.setFont("Helvetica-Bold", 12)
    c.drawString(50, 625, "2. Remote Work & Core Working Hours")
    c.setFont("Helvetica", 10)
    c.drawString(50, 605, "Acme operates on a hybrid-first policy. Employees may work remotely up to 4 days per week.")
    c.drawString(50, 590, "To facilitate real-time collaboration across distributed teams, Acme defines core working hours")
    c.drawString(50, 575, "from 10:00 AM to 4:00 PM Eastern Time (ET), Monday through Friday.")
    c.drawString(50, 560, "All team members must remain available on Slack and email during core working hours.")
    
    c.setFont("Helvetica-Bold", 12)
    c.drawString(50, 530, "3. Parental Leave & Health Benefits")
    c.setFont("Helvetica", 10)
    c.drawString(50, 510, "Acme provides 16 fully paid weeks of parental leave for primary caregivers following childbirth,")
    c.drawString(50, 495, "adoption, or foster placement. Secondary caregivers are eligible for 6 fully paid weeks.")
    c.drawString(50, 480, "Comprehensive medical, dental, and vision insurance coverage begins on day one of employment.")
    
    c.setFont("Helvetica-Bold", 12)
    c.drawString(50, 450, "4. On-Call Compensation Stipend")
    c.setFont("Helvetica", 10)
    c.drawString(50, 430, "Engineering team members assigned to the primary on-call rotation receive a flat $600 weekly")
    c.drawString(50, 415, "standby stipend. Any active incident response time outside normal business hours is compensated")
    c.drawString(50, 400, "at 1.5x the engineer's effective hourly rate.")
    
    c.showPage()
    c.save()
    print(f"Generated {pdf_path}")

def generate_docx():
    docx_path = os.path.join(SAMPLE_DIR, "project_titan_specs.docx")
    doc = docx.Document()
    
    doc.add_heading("Project Titan — Architecture & Technical Specifications", level=1)
    
    doc.add_heading("1. Overview & Objective", level=2)
    doc.add_paragraph(
        "Project Titan is Acme's next-generation real-time analytics and distributed query engine. "
        "The project lead architect is Sarah Jenkins, who also supervises the DevOps release pipeline."
    )
    
    doc.add_heading("2. Core Infrastructure & Databases", level=2)
    doc.add_paragraph(
        "The primary relational persistence layer is PostgreSQL 16, utilizing multi-region streaming "
        "replication with a hot synchronous standby in us-west-2 to guarantee zero data loss (RPO = 0)."
    )
    doc.add_paragraph(
        "Redis 7 is deployed in cluster mode as the distributed caching layer and pub/sub message bus for "
        "session state and query acceleration."
    )
    
    doc.add_heading("3. Performance SLAs & Latency Targets", level=2)
    doc.add_paragraph(
        "Project Titan's search endpoint has a strict latency Service Level Agreement: "
        "the p99 latency must stay under 120 milliseconds under a concurrent throughput of 5,000 requests per second."
    )
    
    doc.add_heading("4. Cloud Budget Allocation", level=2)
    doc.add_paragraph(
        "The approved annual cloud infrastructure budget for Project Titan is $450,000 for the 2024–2025 fiscal year. "
        "This includes AWS compute instances, cross-region replication bandwidth, and reserved Redis clusters."
    )
    
    doc.save(docx_path)
    print(f"Generated {docx_path}")

def generate_markdown():
    md_path = os.path.join(SAMPLE_DIR, "quarterly_financials.md")
    content = """# Acme Corp — Q3 2024 Financial Performance Report

## Executive Summary
In Q3 2024, Acme Corp delivered record growth driven by enterprise customer expansion. 
Total Annual Recurring Revenue (ARR) reached $18.4 million, representing a 42% Year-over-Year (YoY) growth.

## Key Operating Metrics
- **ARR**: $18.4 million (+42% YoY)
- **Net Revenue Retention (NRR)**: 124%
- **Gross Margin**: 78.5%
- **Monthly Net Burn Rate**: $320,000 per month
- **Cash Runway**: 28 months based on current liquid reserves

## Departmental Spending & Cloud Allocations
- Engineering & R&D: $2.4M (including cloud infrastructure allocations)
- Cloud Infrastructure Overall: $1.1M in Q3, within which Project Titan was allocated $112,500 for the quarter (tracking against its $450,000 annual budget).
- Sales & Marketing: $1.8M
- General & Administrative: $650k
"""
    with open(md_path, "w", encoding="utf-8") as f:
        f.write(content)
    print(f"Generated {md_path}")

def generate_text():
    txt_path = os.path.join(SAMPLE_DIR, "incident_runbook.txt")
    content = """ACME CORPORATION — ENGINEERING INCIDENT RUNBOOK
Version: 3.2
Owner & Sign-off Authority: Sarah Jenkins (Lead Architect & DevOps Manager)

1. Incident Severity Definitions
- Sev-1 (Critical): Customer-facing outage or data loss affecting >10% of active sessions.
- Sev-2 (Major): Partial system degradation or high latency without total data unavailability.
- Sev-3 (Minor): Internal tooling issue or cosmetic bug.

2. On-Call Roles & Requirements
- Primary On-Call: First responder, must acknowledge pages within 5 minutes.
- Secondary On-Call: Escalation backup, must acknowledge within 10 minutes.
- All on-call engineers are compensated per the Acme HR on-call stipend policy.

3. Escalation Protocols for Sev-1 Incidents
- 0–5 Minutes: Primary on-call acknowledges page and opens a dedicated incident Slack channel (#incident-active).
- 5–15 Minutes: Incident Commander (IC) assesses blast radius and initiates rollback or circuit breaker.
- 15 Minutes Unresolved: If a Sev-1 incident remains unresolved after 15 minutes, the Incident Commander must page the VP of Engineering and initiate the executive bridge.
- Post-Incident: Root Cause Analysis (RCA) must be published within 48 hours.
"""
    with open(txt_path, "w", encoding="utf-8") as f:
        f.write(content)
    print(f"Generated {txt_path}")

if __name__ == "__main__":
    generate_pdf()
    generate_docx()
    generate_markdown()
    generate_text()
    print("All sample test documents generated successfully!")
