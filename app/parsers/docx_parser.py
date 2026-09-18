"""
DOCX text extraction parser using python-docx.
Extracts both body paragraphs and table cell contents.
"""
import io
import docx

def extract_text_from_docx(file_bytes: bytes, filename: str = "document.docx") -> str:
    """
    Extracts text content from raw DOCX bytes.
    
    Args:
        file_bytes: Raw binary content of the .docx file.
        filename: Original filename for error reporting.
        
    Returns:
        Clean, concatenated text across paragraphs and tables.
        
    Raises:
        ValueError: If the file is empty, corrupted, or contains no extractable text.
    """
    if not file_bytes:
        raise ValueError(f"Uploaded DOCX '{filename}' is empty (0 bytes).")

    try:
        doc = docx.Document(io.BytesIO(file_bytes))
    except Exception as e:
        raise ValueError(f"Failed to parse DOCX '{filename}': {str(e)}")

    extracted_blocks = []

    # 1. Extract paragraphs
    for p in doc.paragraphs:
        text = p.text.strip()
        if text:
            extracted_blocks.append(text)

    # 2. Extract text from tables (structured content)
    for table in doc.tables:
        for row in table.rows:
            row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_texts:
                extracted_blocks.append(" | ".join(row_texts))

    combined_text = "\n\n".join(extracted_blocks).strip()

    if not combined_text:
        raise ValueError(f"Could not extract any text from '{filename}'. The file appears to be blank.")

    return combined_text
